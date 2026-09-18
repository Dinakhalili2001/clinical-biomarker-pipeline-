"""
Plate Reader Raw File Converter
===================================
Converts a raw plate reader export (the full 96-well grid Excel file,
rows A-H, columns 1-12) into the two clean CSV files the ELISA pipeline
expects: standards_data.csv and samples_data.csv.

WHY THIS EXISTS
    The plate reader always exports the ENTIRE plate grid, including
    empty/unused wells. It has no idea which well held which standard
    or which patient sample - only you know that. So this script needs
    a small "well map" telling it what each well actually was.

HOW TO USE (each time you have a new plate)
    1. Save your plate reader export as-is (the raw .xlsx it gives you).
    2. Fill in well_map.csv (see well_map_template.csv for the format) -
       one row per well that actually has data, saying what it was.
    3. Run:
        python convert_plate_reader.py <raw_file.xlsx> <well_map.csv> <plate_id>
    4. This appends the correct rows into standards_data.csv and
       samples_data.csv in ../sample_data/ (creating them if needed),
       ready for the pipeline / Streamlit app to use.
"""

import sys
import re
import pandas as pd
import openpyxl


def read_raw_plate(filepath: str) -> pd.DataFrame:
    """
    Reads a raw plate reader .xlsx export and extracts the 8x12 grid of
    OD values (rows A-H, columns 1-12) into a long-format table:
    columns = well (e.g. "A1"), row, col, raw_OD.
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet = wb.active

    # Find the row containing the "Abs" grid header (column headers 1-12)
    header_row = None
    for row in sheet.iter_rows(min_row=1, max_row=30):
        for cell in row:
            if cell.value == "Abs":
                header_row = cell.row
                break
        if header_row:
            break

    if header_row is None:
        raise ValueError(
            "Could not find the 'Abs' grid header in this file. "
            "This converter expects the standard plate reader export format."
        )

    records = []
    row_letters = ["A", "B", "C", "D", "E", "F", "G", "H"]
    for i, row_letter in enumerate(row_letters):
        data_row = header_row + 1 + i
        for col_num in range(1, 13):
            # Column B in the sheet = well column 1, so offset by 1
            cell = sheet.cell(row=data_row, column=col_num + 1)
            value = cell.value
            if value is None:
                continue
            records.append({
                "well": f"{row_letter}{col_num}",
                "row": row_letter,
                "col": col_num,
                "raw_OD": float(value),
            })

    return pd.DataFrame(records)


def convert(raw_file: str, well_map_file: str, plate_id: str,
            standards_out: str = "../sample_data/standards_data.csv",
            samples_out: str = "../sample_data/samples_data.csv") -> None:
    """
    Reads the raw plate + well map, and appends the correctly labeled
    rows into standards_data.csv and samples_data.csv.
    """
    raw_data = read_raw_plate(raw_file)
    well_map = pd.read_csv(well_map_file)

    merged = well_map.merge(raw_data[["well", "raw_OD"]], on="well", how="left")

    missing = merged[merged["raw_OD"].isna()]
    if not missing.empty:
        print("WARNING: these wells in your well_map were not found in the raw file:")
        print(missing[["well", "role", "label"]])

    merged = merged.dropna(subset=["raw_OD"])
    merged["plate_id"] = plate_id

    # --- Standards ---
    standards_rows = merged[merged["role"] == "standard"][
        ["plate_id", "label", "concentration", "raw_OD"]
    ].rename(columns={"label": "standard"})

    # --- Samples ---
    samples_rows = merged[merged["role"] == "sample"][
        ["plate_id", "label", "timepoint", "raw_OD"]
    ].rename(columns={"label": "patient_id"})

    def append_or_create(df, path):
        try:
            existing = pd.read_csv(path)
            combined = pd.concat([existing, df], ignore_index=True)
        except FileNotFoundError:
            combined = df
        combined.to_csv(path, index=False)

    if not standards_rows.empty:
        append_or_create(standards_rows, standards_out)
        print(f"Added {len(standards_rows)} standard rows to {standards_out}")

    if not samples_rows.empty:
        append_or_create(samples_rows, samples_out)
        print(f"Added {len(samples_rows)} sample rows to {samples_out}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print("Usage: python convert_plate_reader.py <raw_file.xlsx> <well_map.csv> <plate_id>")
        sys.exit(1)

    raw_file, well_map_file, plate_id = sys.argv[1], sys.argv[2], sys.argv[3]
    convert(raw_file, well_map_file, plate_id)
